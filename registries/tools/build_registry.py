"""Regenerate the AIDCF-007A registry v1.1 outputs from registries/controls/controls.json:
controls.csv, gate-mappings.csv, evidence-profiles.csv, parameters.csv, the criticalGates column of
the domain and family registries, and the v1.1 workbook. Requires openpyxl. Row status is the
AIDCF-014 record lifecycle code CANDIDATE (proposed, awaiting review) until amendment AMD-002 is ratified."""
import json,glob,csv,collections,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
import os
R=os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)),'..','..'))+'/'
# Source of truth: the full control definitions in registries/controls/controls.json
C=json.load(open(R+'registries/controls/controls.json',encoding='utf-8'))['controls']
ORDER="GOV CAP DEP PHY POW THM ENV CMP STO NET OPS RES SUP SEC EVA".split()
C.sort(key=lambda c:(ORDER.index(c['controlId'][4:7]),c['controlId']))
assert len(C)==360
doms=list(csv.DictReader(open(R+'registries/domains/domains.csv')))
fams=list(csv.DictReader(open(R+'registries/families/control-families.csv')))
famname={f['familyId']:f['familyName'] for f in fams}
domname={d['domainId']:d['domainName'] for d in doms}
j=lambda x:'; '.join(x)
def crit(c,role): return [g['gateId'] for g in c['gateMappings'] if g['role']==role]
# full JSON
json.dump({"registry":"AIDCF-007A Canonical Master Control Registry","version":"1.1","status":"PROPOSED - Draft for controlled review","amendmentId":"AIDCF-007A-AMD-002","date":"2026-10-10","controls":C},open(R+'registries/controls/controls.json','w'),indent=1,ensure_ascii=False)
# controls.csv
cols="controlId,familyId,domainId,title,objective,requirement,claimType,requiredAssuranceLevel,controlType,profileTier,sufficiencyFloor,criticalGates,supportingGates,status,version,amendmentId".split(',')
with open(R+'registries/controls/controls.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(cols)
    for c in C:
        ep=c['evidenceProfile']
        w.writerow([c['controlId'],c['familyId'],c['controlId'][4:7],c['title'],c['objective'],c['requirement'],ep['claimType'],c['requiredAssuranceLevel'],c['controlType'],ep['profileTier'],ep['sufficiencyFloor'],';'.join(crit(c,'CRITICAL')),';'.join(crit(c,'SUPPORTING')),'CANDIDATE','1.1','AIDCF-007A-AMD-002'])
with open(R+'registries/controls/gate-mappings.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['controlId','gateId','role'])
    for c in C:
        for g in c['gateMappings']: w.writerow([c['controlId'],g['gateId'],g['role']])
with open(R+'registries/controls/evidence-profiles.csv','w',newline='') as f:
    k="claimType profileTier requiredClasses optionalClasses qualityFloor sufficiencyFloor stateCoverage periodCoverage independence mandatoryItems acceptableAlternatives unknownEffect reassessmentTriggers".split()
    w=csv.writer(f); w.writerow(['controlId']+k)
    for c in C:
        ep=c['evidenceProfile']; w.writerow([c['controlId']]+[j(ep[x]) if isinstance(ep[x],list) else ep[x] for x in k])
with open(R+'registries/controls/parameters.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['controlId','name','unit','illustrativeDefault','basis'])
    for c in C:
        for p in c['parameters']: w.writerow([c['controlId'],p.get('name'),p.get('unit'),p.get('illustrativeDefault'),p.get('basis')])
# family/domain derived gates
fg=collections.defaultdict(set); dg=collections.defaultdict(set)
for c in C:
    for g in crit(c,'CRITICAL'): fg[c['familyId']].add(g); dg[c['controlId'][4:7]].add(g)
for rows,key,agg,p in [(fams,'familyId',fg,'registries/families/control-families.csv'),(doms,'domainId',dg,'registries/domains/domains.csv')]:
    hdr=list(rows[0].keys())
    with open(R+p,'w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=hdr); w.writeheader()
        for r in rows:
            r['criticalGates']=';'.join(sorted(agg[r[key]])); r['version']='1.1'; r['status']='CANDIDATE'; w.writerow(r)
# workbook
wb=openpyxl.Workbook(); H=Font(bold=True,color='FFFFFF'); F=PatternFill('solid',fgColor='17365D')
def sheet(name,header,rows,widths=None):
    ws=wb.create_sheet(name); ws.append(header)
    for c in ws[1]: c.font=H; c.fill=F; c.alignment=Alignment(wrap_text=True,vertical='top')
    for r in rows: ws.append(r)
    ws.freeze_panes='B2'
    for i,_ in enumerate(header):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=(widths or {}).get(i,28)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    return ws
wb.remove(wb.active)
gates=list(csv.DictReader(open(R+'registries/gates/critical-gates.csv')))
crc=collections.Counter(g for c in C for g in crit(c,'CRITICAL')); spc=collections.Counter(g for c in C for g in crit(c,'SUPPORTING'))
ws=sheet('README',['Topic','Guidance'],[
 ['Registry','AIDCF-007A Canonical Master Control Registry, version 1.1 (amendment AIDCF-007A-AMD-002), 2026-10-10.'],
 ['Status','PROPOSED - Draft for controlled review. Version 1.0 (AIDCF-007A-AMD-001) remains the approved cardinality and identifier baseline until v1.1 completes controlled methodology review and independent ratification.'],
 ['What changed in v1.1','Control content rewritten per control (objective, requirement, criteria, parameters, evidence profile, procedure, testing safety, UNKNOWN rule); IDs and 15/90/360 structure unchanged; duplicate titles fixed (MCF-NET-004, MCF-STO-019); control-specific gate mappings with CRITICAL/SUPPORTING roles replace whole-domain mapping; Required Assurance Level assigned per control from its claim type (AL2-AL5).'],
 ['Gate roll-up','A gate cannot PASS unless every in-scope CRITICAL control is CONFORMING at its required claim type; PARTIAL never yields PASS. SUPPORTING controls inform the gate evaluation. CG-11 and CG-12 are meta-gates evaluated across all controls. See AIDCF-010.'],
 ['Vocabularies','All codes are canonical per AIDCF-002 and registries/vocabularies/.'],
 ['Constitutional rule','Controls do not create assurance by existence. Missing evidence is UNKNOWN, not N/A. Scores never override gates (AIDCF-006).'],
 ['Parameters','Engineering values are capability-profile parameters; defaults are illustrative unless attributed to a named source.'],
 ['Machine-readable copies','registries/controls/controls.json (full), controls.csv, gate-mappings.csv, evidence-profiles.csv, parameters.csv.']],{0:24,1:120})
sheet('Summary',['Metric','Value'],[['Domains',len(doms)],['Families',len(fams)],['Controls',len(C)],['Controls per family',4],['Unique titles',len({c['title'] for c in C})],['CRITICAL gate mappings',sum(crc.values())],['SUPPORTING gate mappings',sum(spc.values())]]+[[f'Claim type {k}',v] for k,v in collections.Counter(c['evidenceProfile']['claimType'] for c in C).items()],{0:30,1:20})
sheet('Domains',['Domain Code','Domain','Families','Controls','Gates with CRITICAL controls'],[[d['domainId'],d['domainName'],6,24,'; '.join(sorted(dg[d['domainId']]))] for d in doms],{1:55})
sheet('Families',['Family ID','Domain Code','Control Family','Controls','Gates with CRITICAL controls'],[[f['familyId'],f['domainId'],f['familyName'],'; '.join(c['controlId']+' '+c['title'] for c in C if c['familyId']==f['familyId']),'; '.join(sorted(fg[f['familyId']]))] for f in fams],{2:40,3:90})
hdr=['Control ID','Domain Code','Domain','Family ID','Control Family','Control Title','Control Objective','Normative Requirement','Assessment Criteria','Parameters','Minimum Evidence','Assessment Methods','Assessment Procedure','Control Type','Claim Type','Required Assurance Level','Lifecycle Stages','Operating States','Applicability','CRITICAL Gates','SUPPORTING Gates','Failure Modes','Ontology Classes','Testing Safety','UNKNOWN Rule','Informative References','Related Controls','Status','Version','Amendment ID']
num=lambda xs:'\n'.join(f'{i+1}) {x}' for i,x in enumerate(xs))
rows=[[c['controlId'],c['controlId'][4:7],domname[c['controlId'][4:7]],c['familyId'],famname[c['familyId']],c['title'],c['objective'],c['requirement'],num(c['assessmentCriteria']),'\n'.join(f"{p.get('name')} [{p.get('unit')}] default {p.get('illustrativeDefault')} ({p.get('basis')})" for p in c['parameters']),j(c['minimumEvidence']),j(c['assessmentMethods']),num(c['assessmentProcedure']),c['controlType'],c['evidenceProfile']['claimType'],c['requiredAssuranceLevel'],j(c['lifecycleStages']),j(c['operatingStates']),c['applicability'],j(crit(c,'CRITICAL')),j(crit(c,'SUPPORTING')),j(c['failureModes']),j(c['ontologyClasses']),c['testingSafety'],c['unknownRule'],j(c['informativeReferences']),j(c['relatedControls']),'CANDIDATE','1.1','AIDCF-007A-AMD-002'] for c in C]
sheet('Control Registry',hdr,rows,{6:50,7:70,8:70,12:70})
k="claimType profileTier requiredClasses optionalClasses qualityFloor sufficiencyFloor stateCoverage periodCoverage independence mandatoryItems acceptableAlternatives unknownEffect reassessmentTriggers".split()
sheet('Evidence Profiles',['Control ID']+k,[[c['controlId']]+[j(c['evidenceProfile'][x]) if isinstance(c['evidenceProfile'][x],list) else c['evidenceProfile'][x] for x in k] for c in C])
sheet('Gate Mapping',['Control ID','Control Title','Gate','Role'],[[c['controlId'],c['title'],g['gateId'],g['role']] for c in C for g in c['gateMappings']])
sheet('Critical Gates',['Gate','Name','Scope','Evidence Floor','CRITICAL controls','SUPPORTING controls','CRITICAL control IDs'],[[g['gateId'],g['name'],g['scope'],g['evidenceFloor'],crc[g['gateId']],spc[g['gateId']],'; '.join(c['controlId'] for c in C if g['gateId'] in crit(c,'CRITICAL')) or ('Meta-gate: evaluated across all controls' if g['gateId'] in ('CG-11','CG-12') else '')] for g in gates],{2:60,6:80})
sheet('Parameters',['Control ID','Parameter','Unit','Illustrative Default','Basis'],[[c['controlId'],p.get('name'),p.get('unit'),p.get('illustrativeDefault'),p.get('basis')] for c in C for p in c['parameters']])
sheet('Assessment Template',['Control ID','Control Title','Claim Type','Required AL','In Scope (Y/N)','Result (CONFORMING/PARTIAL/NONCONFORMING/UNKNOWN/NOT_APPLICABLE/NOT_ASSESSED)','Evidence Package ID','Quality Vector','Sufficiency (S0-S5)','Confidence (CONF-0..4)','UNKNOWN Type/Severity','Findings','CRITICAL Gates Affected','Assessor','Reviewer','Date'],[[c['controlId'],c['title'],c['evidenceProfile']['claimType'],c['requiredAssuranceLevel'],'','NOT_ASSESSED','','','','','','',j(crit(c,'CRITICAL')),'','',''] for c in C])
sheet('Amendment Log',['Amendment ID','Date','Status','Summary'],[['AIDCF-007A-AMD-001','(as stated in v1.0 issue)','Approved by baseline owner; independent ratification pending','Adopted the 15/90/360 cardinality and MCF identifier baseline and deprecated narrative statements of 14/84/336.'],['AIDCF-007A-AMD-002','2026-10-10','PROPOSED - Draft for controlled review','Per-control content rewrite, duplicate-title fixes, control-specific CRITICAL/SUPPORTING gate mapping, per-control assurance levels, evidence requirement profiles, assessment template restored.']],{3:110})
out=R+'standards/aidcf-007a/registry/AIDCF-007A_Canonical_Master_Control_Registry_15-90-360_v1.1.xlsx'
wb.save(out); print('saved',out)
